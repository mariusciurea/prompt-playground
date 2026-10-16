"""Parse e-mail lists from pasted text, CSV, or XLSX files into a validation report."""

import csv
import io
import re
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from typing import Any

from email_validator import EmailNotValidError, validate_email
from openpyxl import load_workbook

from app.core.exceptions import AppError

MAX_IMPORT_BYTES = 1_000_000
MAX_ROWS = 2000
_MAX_CELLS_PER_ROW = 20
_ANGLE_ADDRESS = re.compile(r"<([^<>\s]+@[^<>\s]+)>")
_TEXT_SEPARATORS = re.compile(r"[\s,;]+")


class ImportFileError(AppError):
    """The uploaded file cannot be read."""

    code = "invalid_import"


@dataclass(frozen=True, slots=True)
class InvalidEntry:
    """A value that is not a usable e-mail address."""

    row: int
    value: str
    reason: str


@dataclass(slots=True)
class ParsedEmails:
    """Unique valid addresses in input order, plus everything that was rejected."""

    valid: list[str] = field(default_factory=list)
    invalid: list[InvalidEntry] = field(default_factory=list)
    duplicates: list[str] = field(default_factory=list)


def normalize_email(value: str) -> str:
    """Validate an address (syntax only) and return it lower-cased."""
    return validate_email(value, check_deliverability=False).normalized.lower()


def _candidate(cell: Any) -> str | None:
    if cell is None:
        return None
    text = str(cell).strip()
    if "@" not in text:
        return None
    match = _ANGLE_ADDRESS.search(text)
    return match.group(1) if match else text


def _rows_from_csv(content: bytes) -> Iterator[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            text[:4096], delimiters=",;\t"
        )
    except csv.Error:
        dialect = csv.excel
    yield from csv.reader(io.StringIO(text), dialect)


def _rows_from_xlsx(content: bytes) -> Iterator[tuple[Any, ...]]:
    try:
        # data_only: read cached cell values, never formulas; defusedxml guards the XML parsing.
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError("The Excel file could not be read.") from exc
    try:
        sheet = workbook.worksheets[0]
        yield from sheet.iter_rows(
            max_row=MAX_ROWS + 1, max_col=_MAX_CELLS_PER_ROW, values_only=True
        )
    finally:
        workbook.close()


def _collect(rows: Iterable[Iterable[Any]], *, header_allowed: bool = True) -> ParsedEmails:
    result = ParsedEmails()
    seen: set[str] = set()
    for number, row in enumerate(rows, start=1):
        if number > MAX_ROWS:
            raise ImportFileError(f"Too many rows; the limit is {MAX_ROWS}.")
        cells = list(row)[:_MAX_CELLS_PER_ROW]
        value = next((c for c in map(_candidate, cells) if c), None)
        if value is None:
            is_header = header_allowed and number == 1
            if not is_header and any(str(c).strip() for c in cells if c is not None):
                result.invalid.append(InvalidEntry(number, _preview(cells), "no e-mail address"))
            continue  # blank rows and a header row are skipped
        try:
            email = normalize_email(value)
        except EmailNotValidError as exc:
            result.invalid.append(InvalidEntry(number, value[:254], str(exc)))
            continue
        if email in seen:
            result.duplicates.append(email)
        else:
            seen.add(email)
            result.valid.append(email)
    return result


def _preview(cells: list[Any]) -> str:
    return ", ".join(str(c) for c in cells if c is not None)[:254]


def parse_text(text: str) -> ParsedEmails:
    """Parse addresses separated by new lines, commas, semicolons, or spaces."""
    tokens = [t for t in _TEXT_SEPARATORS.split(text) if t]
    return _collect(([token] for token in tokens), header_allowed=False)


def parse_file(filename: str, content: bytes) -> ParsedEmails:
    """Parse a CSV/TXT or XLSX upload; the first e-mail-like cell of each row is used."""
    if len(content) > MAX_IMPORT_BYTES:
        raise ImportFileError("The file is too large; the limit is 1 MB.")
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension in {"csv", "txt"}:
        return _collect(_rows_from_csv(content))
    if extension == "xlsx":
        return _collect(_rows_from_xlsx(content))
    raise ImportFileError("Upload a .csv, .txt, or .xlsx file.")
